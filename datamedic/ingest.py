"""Source-file validation and reading.

Every format is reduced to ONE internal reader: a stream of string-typed
pandas chunks with positional integer column labels. CSV files are read
directly from the preserved original; Excel sheets are streamed into an
intermediate CSV first (the original workbook is never modified).
"""
from __future__ import annotations

import codecs
import csv
import datetime as dt
import logging
import os
import time
import zipfile
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Iterator

import pandas as pd

from . import config
from .errors import DataMedicError

log = logging.getLogger("datamedic.ingest")

XLS_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ZIP_MAGIC = b"PK\x03\x04"


@dataclass
class SourceInfo:
    kind: str                       # csv | xlsx | xls
    read_path: str                  # file the chunk reader streams from (relative to job dir)
    encoding: str = "utf-8"
    delimiter: str = ","
    sheet: str | None = None
    sheets: list[str] = field(default_factory=list)
    header: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    def to_dict(self):
        return asdict(self)


def chunk_rows_for(ncols: int) -> int:
    rows = config.TARGET_CELLS_PER_CHUNK // max(1, ncols)
    return max(config.MIN_CHUNK_ROWS, min(config.MAX_CHUNK_ROWS, rows))


# --------------------------------------------------------------------------
# Quick validation (cheap; runs inside the upload request)
# --------------------------------------------------------------------------
def validate_upload(path: Path, ext: str) -> None:
    size = path.stat().st_size
    if size == 0:
        raise DataMedicError("The file is empty (0 bytes).", "empty_file", 400)
    if size > config.MAX_FILE_BYTES:
        raise DataMedicError(too_large_message(size), "file_too_large", 413)
    with open(path, "rb") as f:
        head = f.read(65536)
    if ext == ".xlsx":
        if not head.startswith(ZIP_MAGIC):
            raise DataMedicError(
                "This file is named .xlsx but is not a valid Excel workbook "
                "(it may be corrupted or a renamed file).", "bad_format", 422)
    elif ext == ".xls":
        if not head.startswith(XLS_MAGIC):
            raise DataMedicError(
                "This file is named .xls but is not a valid legacy Excel workbook. "
                "If it is really .xlsx or .csv, rename it correctly.", "bad_format", 422)
    else:
        if head.startswith(ZIP_MAGIC) or head.startswith(XLS_MAGIC):
            raise DataMedicError(
                "This file is named .csv but looks like an Excel workbook. "
                "Rename it to .xlsx / .xls.", "bad_format", 422)
        if head.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
            return
        if b"\x00" in head:
            raise DataMedicError(
                "This file contains binary data and is not a text CSV file.", "bad_format", 422)


def too_large_message(size: int) -> str:
    return (f"This file is {size:,} bytes ({size / 1024 ** 2:,.1f} MB), which exceeds the "
            f"{config.MAX_FILE_BYTES // 1024 ** 2} MB "
            f"({config.MAX_FILE_BYTES:,} bytes) limit.")


# --------------------------------------------------------------------------
# CSV detection
# --------------------------------------------------------------------------
def detect_encoding(path: Path) -> tuple[str, list[str]]:
    with open(path, "rb") as f:
        head = f.read(4)
    if head.startswith(codecs.BOM_UTF8):
        return "utf-8-sig", []
    if head.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return "utf-16", []

    def strict_ok(enc: str) -> bool:
        dec = codecs.getincrementaldecoder(enc)("strict")
        try:
            with open(path, "rb") as fh:
                while True:
                    blk = fh.read(1 << 20)
                    if not blk:
                        break
                    dec.decode(blk)
            dec.decode(b"", final=True)
            return True
        except UnicodeDecodeError:
            return False

    if strict_ok("utf-8"):
        return "utf-8", []
    if strict_ok("cp1252"):
        return "cp1252", ["The file is not valid UTF-8; it was read as Windows-1252 (ANSI). "
                          "Check that accented characters look right in the preview."]
    return "latin-1", ["The file uses an unknown text encoding; it was read as Latin-1. "
                       "Some characters may be wrong."]


def detect_delimiter(path: Path, encoding: str) -> str:
    with open(path, "r", encoding=encoding, errors="replace", newline="") as f:
        sample = f.read(65536)
    lines = sample.splitlines()[:60]
    best, best_score = ",", (0.0, 0)
    for delim in (",", ";", "\t", "|"):
        counts = [len(r) for r in csv.reader(lines, delimiter=delim) if r]
        if not counts:
            continue
        mode = max(set(counts), key=counts.count)
        if mode < 2:
            continue
        score = (counts.count(mode) / len(counts), mode)
        if score > best_score:
            best, best_score = delim, score
    return best


def read_csv_chunks(job_dir: Path, src: SourceInfo, chunksize: int) -> Iterator[pd.DataFrame]:
    """Yield string chunks (positional int columns, 0-based RangeIndex per chunk),
    EXCLUDING the header row. Raises DataMedicError on malformed input."""
    path = job_dir / src.read_path
    try:
        reader = pd.read_csv(
            path, sep=src.delimiter, encoding=src.encoding, header=None, dtype=str,
            keep_default_na=False, na_filter=False, chunksize=chunksize,
            on_bad_lines="error", skip_blank_lines=True, engine="c")
        first = True
        for chunk in reader:
            if first:
                chunk = chunk.iloc[1:]
                first = False
            if len(chunk):
                yield chunk.reset_index(drop=True)
    except pd.errors.ParserError as exc:
        raise DataMedicError(
            "The CSV is malformed: " + _friendly_parser_error(str(exc)) +
            " Fix the file (quotes, delimiters) and upload it again.", "malformed_csv", 422)
    except pd.errors.EmptyDataError:
        raise DataMedicError("The file contains no data.", "empty_data", 422)
    except UnicodeDecodeError:
        raise DataMedicError("The file's text encoding could not be decoded.", "encoding", 422)


def _friendly_parser_error(msg: str) -> str:
    msg = msg.replace("Error tokenizing data. C error: ", "").strip()
    return msg[:200] if msg else "rows have inconsistent numbers of columns."


def prepare_csv(job_dir: Path, source_rel: str) -> SourceInfo:
    path = job_dir / source_rel
    enc, warns = detect_encoding(path)
    delim = detect_delimiter(path, enc)
    src = SourceInfo(kind="csv", read_path=source_rel, encoding=enc, delimiter=delim, warnings=warns)
    try:
        first = pd.read_csv(path, sep=delim, encoding=enc, header=None, dtype=str,
                            keep_default_na=False, na_filter=False, nrows=1)
    except pd.errors.EmptyDataError:
        raise DataMedicError("The file contains no data (only blank lines).", "empty_data", 422)
    except pd.errors.ParserError as exc:
        raise DataMedicError("The CSV header could not be parsed: " + _friendly_parser_error(str(exc)),
                             "malformed_csv", 422)
    src.header = [str(v) for v in first.iloc[0].tolist()]
    check_width(len(src.header))
    return src


def check_width(ncols: int) -> None:
    if ncols > config.MAX_COLUMNS:
        raise DataMedicError(
            f"The file has {ncols:,} columns; the maximum supported is {config.MAX_COLUMNS:,}.",
            "too_many_columns", 422)


# --------------------------------------------------------------------------
# Excel -> intermediate CSV
# --------------------------------------------------------------------------
def _fmt_number(v) -> str:
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v != v or v in (float("inf"), float("-inf")):
            return ""
        if v.is_integer() and abs(v) < 1e15:
            return str(int(v))
        return repr(v)
    return str(v)


def _fmt_cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d") if (v.hour, v.minute, v.second, v.microsecond) == (0, 0, 0, 0) \
            else v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, dt.time):
        return v.strftime("%H:%M:%S")
    if isinstance(v, dt.timedelta):
        return str(v)
    if isinstance(v, (bool, int, float)):
        return _fmt_number(v)
    return str(v)


def _check_xlsx_zip(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as zf:
            infos = zf.infolist()
            total = sum(i.file_size for i in infos)
            comp = sum(i.compress_size for i in infos) or 1
            names = {i.filename for i in infos}
            if "xl/workbook.xml" not in names:
                raise DataMedicError("This is not a valid .xlsx workbook (workbook.xml is missing).",
                                     "bad_format", 422)
            if "xl/encryptedPackage" in names or "EncryptedPackage" in names:
                raise DataMedicError("Password-protected workbooks are not supported.", "encrypted", 422)
    except zipfile.BadZipFile:
        raise DataMedicError("The .xlsx file is corrupted (invalid zip container).", "bad_format", 422)
    if total > config.MAX_XLSX_UNCOMPRESSED_BYTES:
        raise DataMedicError(
            f"The workbook expands to {total / 1024 ** 3:.1f} GB of XML, above the "
            f"{config.MAX_XLSX_UNCOMPRESSED_BYTES // 1024 ** 3} GB safety limit.", "xlsx_too_large", 413)
    if total / comp > config.MAX_XLSX_COMPRESSION_RATIO and total > 100 * 1024 ** 2:
        raise DataMedicError("The workbook has a suspiciously high compression ratio and was rejected.",
                             "xlsx_suspicious", 422)


def _iter_xlsx_rows(path: Path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            yield ws.title, ws.iter_rows(values_only=True)
    finally:
        wb.close()


def _iter_xls_rows(path: Path):
    import xlrd
    try:
        book = xlrd.open_workbook(str(path), on_demand=True)
    except xlrd.XLRDError as exc:
        msg = str(exc).lower()
        if "encrypt" in msg or "password" in msg:
            raise DataMedicError("Password-protected workbooks are not supported.", "encrypted", 422)
        raise DataMedicError("The .xls file could not be read (it may be corrupted).", "bad_format", 422)

    def rows(sheet):
        for r in range(sheet.nrows):
            out = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                t, v = cell.ctype, cell.value
                if t == xlrd.XL_CELL_EMPTY or t == xlrd.XL_CELL_BLANK:
                    out.append(None)
                elif t == xlrd.XL_CELL_DATE:
                    try:
                        out.append(xlrd.xldate.xldate_as_datetime(v, book.datemode))
                    except Exception:
                        out.append(v)
                elif t == xlrd.XL_CELL_BOOLEAN:
                    out.append(bool(v))
                elif t == xlrd.XL_CELL_ERROR:
                    out.append(xlrd.error_text_from_code.get(v, "#ERROR"))
                else:
                    out.append(v)
            yield tuple(out)

    try:
        for i in range(book.nsheets):
            sh = book.sheet_by_index(i)
            yield sh.name, rows(sh)
    finally:
        book.release_resources()


def _sheet_names(path: Path, kind: str) -> list[str]:
    try:
        if kind == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True)
            try:
                return list(wb.sheetnames)
            finally:
                wb.close()
        import xlrd
        book = xlrd.open_workbook(str(path), on_demand=True)
        try:
            return list(book.sheet_names())
        finally:
            book.release_resources()
    except Exception:
        return []


def prepare_excel(job_dir: Path, source_rel: str, kind: str) -> SourceInfo:
    """Stream the first non-empty sheet into work/ingested.csv (UTF-8)."""
    path = job_dir / source_rel
    work = job_dir / "work"
    work.mkdir(exist_ok=True)
    out_rel = "work/ingested.csv"
    out_path = job_dir / out_rel
    if kind == "xlsx":
        _check_xlsx_zip(path)
    else:
        if path.stat().st_size > 100 * 1024 ** 2:
            raise DataMedicError(
                "Legacy .xls files over 100 MB cannot be read safely. Save it as .xlsx or .csv.",
                "xls_too_large", 413)
    src = SourceInfo(kind=kind, read_path=out_rel, encoding="utf-8", delimiter=",")
    started = time.monotonic()
    try:
        source_iter = _iter_xlsx_rows(path) if kind == "xlsx" else _iter_xls_rows(path)
        chosen = None
        for title, rows in source_iter:
            src.sheets.append(title)
            if chosen is not None:
                continue            # keep listing sheet names only
            # pass 1: widest non-empty cell + first row
            width, nrows = 0, 0
            for row in rows:
                last = 0
                for i in range(len(row) - 1, -1, -1):
                    if row[i] not in (None, ""):
                        last = i + 1
                        break
                if last:
                    width = max(width, last)
                    nrows += 1
                    if nrows % 2000 == 0 and time.monotonic() - started > config.EXCEL_CONVERT_TIMEOUT_S:
                        raise DataMedicError("Reading the workbook took too long and was stopped.",
                                             "timeout", 504)
            if not nrows:
                continue
            check_width(width)
            chosen = title
            src.sheet = title
            # pass 2 re-opens the sheet to write rows
            for t2, rows2 in (_iter_xlsx_rows(path) if kind == "xlsx" else _iter_xls_rows(path)):
                if t2 != title:
                    continue
                import csv as _csv
                with open(out_path, "w", encoding="utf-8", newline="") as fh:
                    w = _csv.writer(fh)
                    n = 0
                    for row in rows2:
                        vals = [_fmt_cell(v) for v in row[:width]]
                        if len(vals) < width:
                            vals += [""] * (width - len(vals))
                        if not any(v != "" for v in vals):
                            continue
                        w.writerow(vals)
                        n += 1
                        if n % 5000 == 0 and time.monotonic() - started > config.EXCEL_CONVERT_TIMEOUT_S:
                            raise DataMedicError("Reading the workbook took too long and was stopped.",
                                                 "timeout", 504)
                break
            break
        if chosen is None:
            raise DataMedicError("The workbook contains no data in any sheet.", "empty_data", 422)
    except DataMedicError:
        raise
    except MemoryError:
        raise DataMedicError("The workbook is too large to read with the available memory.",
                             "memory", 413)
    except Exception as exc:  # corrupted / unsupported workbook
        log.warning("Excel read failure: %r", exc)
        raise DataMedicError("The workbook could not be read. It may be corrupted, encrypted "
                             "or use unsupported features.", "bad_format", 422)
    src.sheets = _sheet_names(path, kind) or src.sheets
    others = [s for s in src.sheets if s != src.sheet]
    if others:
        src.warnings.append(
            f"Only the first sheet with data (\"{src.sheet}\") was analysed. "
            f"Other sheets ignored: {', '.join(others[:8])}{'…' if len(others) > 8 else ''}.")
    src.warnings.append("Excel formulas are read as their last saved values; fully blank rows were skipped.")
    first = pd.read_csv(out_path, header=None, dtype=str, keep_default_na=False, na_filter=False, nrows=1)
    src.header = [str(v) for v in first.iloc[0].tolist()]
    return src


def prepare_source(job_dir: Path, source_rel: str, kind: str) -> SourceInfo:
    if kind == "csv":
        return prepare_csv(job_dir, source_rel)
    return prepare_excel(job_dir, source_rel, kind)
