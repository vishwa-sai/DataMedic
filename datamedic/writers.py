"""Streaming output writers (CSV, XLSX) plus post-write verification."""
from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import pandas as pd

from .errors import DataMedicError

_INT_RE = re.compile(r"-?[0-9]{1,15}")


class CsvWriter:
    def __init__(self, path: Path, names: list[str], kept: list[int]):
        self.kept = kept
        self.fh = open(path, "w", encoding="utf-8-sig", newline="")
        pd.DataFrame([names]).to_csv(self.fh, header=False, index=False, lineterminator="\r\n")
        self.warnings: list[str] = []

    def write(self, df: pd.DataFrame) -> None:
        df[self.kept].to_csv(self.fh, header=False, index=False, lineterminator="\r\n")

    def close(self) -> None:
        self.fh.close()

    def abort(self) -> None:
        self.fh.close()


class XlsxWriter:
    """Write-only (streaming) workbook. Numeric/date columns are typed."""

    def __init__(self, path: Path, names: list[str], kept: list[int], numeric: set[int], dates: dict[int, bool]):
        import openpyxl
        from openpyxl.cell.cell import WriteOnlyCell, ILLEGAL_CHARACTERS_RE
        self._cell, self._illegal = WriteOnlyCell, ILLEGAL_CHARACTERS_RE
        self.path, self.kept = path, kept
        self.wb = openpyxl.Workbook(write_only=True)
        self.ws = self.wb.create_sheet("Cleaned")
        self.kinds = ["date" if (c in dates and not dates[c]) else "datetime" if c in dates
                      else "num" if c in numeric else "text" for c in kept]
        self.removed_ctrl = 0
        self.ws.append([self._text(n) for n in names])
        self.warnings: list[str] = []

    def _text(self, v: str):
        if v == "":
            return None
        cleaned, n = self._illegal.subn("", v)
        self.removed_ctrl += n
        cell = self._cell(self.ws, value=cleaned)
        cell.data_type = "s"       # never let Excel interpret text as a formula
        return cell

    def _conv(self, kind: str, v: str):
        if v == "":
            return None
        if kind == "num":
            try:
                return int(v) if _INT_RE.fullmatch(v) else float(v)
            except ValueError:
                return self._text(v)
        if kind == "date":
            try:
                return dt.date.fromisoformat(v)
            except ValueError:
                return self._text(v)
        if kind == "datetime":
            try:
                return dt.datetime.fromisoformat(v)
            except ValueError:
                return self._text(v)
        return self._text(v)

    def write(self, df: pd.DataFrame) -> None:
        cols = [df[c].tolist() for c in self.kept]
        kinds = self.kinds
        for row in zip(*cols):
            self.ws.append([self._conv(k, v) for k, v in zip(kinds, row)])

    def abort(self) -> None:
        try:
            self.wb.close()
        except Exception:
            pass

    def close(self) -> None:
        self.wb.save(self.path)
        self.wb.close()
        if self.removed_ctrl:
            self.warnings.append(f"{self.removed_ctrl} unsupported control character(s) were removed for Excel compatibility.")


def verify_output(path: Path, fmt: str, expected_rows: int, expected_cols: int) -> None:
    """Re-open the finished file and confirm its shape. Raises on mismatch."""
    if not path.exists() or path.stat().st_size == 0:
        raise DataMedicError("The output file was not created.", "verify_failed", 500)
    if fmt == "csv":
        rows, cols = 0, None
        for ch in pd.read_csv(path, header=None, dtype=str, keep_default_na=False, na_filter=False,
                              skip_blank_lines=False, encoding="utf-8-sig", chunksize=200_000):
            rows += len(ch)
            cols = ch.shape[1]
        rows -= 1                                   # header
        if cols is None:
            raise DataMedicError("The output file is empty.", "verify_failed", 500)
    else:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            ws = wb["Cleaned"]
            rows, cols = -1, 0
            for row in ws.iter_rows(values_only=True):
                rows += 1
                cols = max(cols, len(row))
        finally:
            wb.close()
    if rows != expected_rows or cols != expected_cols:
        raise DataMedicError(
            f"Output verification failed (expected {expected_rows:,} rows × {expected_cols} columns, "
            f"found {rows:,} × {cols}). The result was discarded.", "verify_failed", 500)
